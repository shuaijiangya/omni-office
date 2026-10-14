#!/usr/bin/env python3
"""Generate a reproducible, intentionally incomplete DOCX for completion testing."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def set_font(run, *, chinese="宋体", western="Times New Roman", size=10.5,
             bold=False, color=None):
    run.font.name = western
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), chinese)
    fonts.set(qn("w:ascii"), western)
    fonts.set(qn("w:hAnsi"), western)


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tblpr = table._tbl.tblPr
    borders = tblpr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn("w:" + side))
        if edge is None:
            edge = OxmlElement("w:" + side)
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "5")
        edge.set(qn("w:color"), "D9D9D9")


def set_table_width(table):
    tblpr = table._tbl.tblPr
    width = tblpr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tblpr.insert(0, width)
    # 5000 means 100% of the available text width in OOXML.
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), "5000")
    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")


def remove_paragraph_border(element):
    ppr = element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    trpr.append(flag)


def set_cell_text(cell, text, *, header=False, center=False):
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            set_font(run, chinese="黑体" if header else "宋体", size=9.5,
                     bold=False, color=(0, 0, 0))


def add_table(document, headers, rows, widths, *, title=None):
    if title:
        caption = document.add_paragraph(style="Caption")
        caption.paragraph_format.keep_with_next = True
        caption.add_run(title)
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    set_table_borders(table)
    for i, (label, width) in enumerate(zip(headers, widths)):
        table.columns[i].width = Inches(width)
        cell = table.rows[0].cells[i]
        cell.width = Inches(width)
        set_cell_shading(cell, "DDEAF4")
        set_cell_text(cell, label, header=True)
    repeat_header(table.rows[0])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for i, (value, width) in enumerate(zip(values, widths)):
            cell = row.cells[i]
            cell.width = Inches(width)
            if row_index % 2:
                set_cell_shading(cell, "F7FAFC")
            set_cell_text(cell, value, center=(i == 0 and len(headers) <= 3))
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_para(document, text, *, style=None):
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)
    return paragraph


def add_styled_placeholder(document, label):
    paragraph = document.add_paragraph()
    label_run = paragraph.add_run(label)
    set_font(label_run, color=(0, 0, 0))
    value_run = paragraph.add_run("待填写")
    set_font(value_run, bold=True, color=(192, 0, 0))
    return paragraph


def set_style_font(style, chinese, western, size, bold=False):
    style.font.name = western
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), chinese)
    fonts.set(qn("w:ascii"), western)
    fonts.set(qn("w:hAnsi"), western)


def build(output):
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    set_style_font(normal, "宋体", "Times New Roman", 10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.2
    title_style = document.styles["Title"]
    set_style_font(title_style, "黑体", "Times New Roman", 18)
    remove_paragraph_border(title_style.element)
    title_style.paragraph_format.space_after = Pt(13)
    for name, size in (("Heading 1", 13), ("Heading 2", 11)):
        style = document.styles[name]
        set_style_font(style, "黑体", "Times New Roman", size)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    caption_style = document.styles["Caption"]
    set_style_font(caption_style, "宋体", "Times New Roman", 9.5)
    caption_style.paragraph_format.space_before = Pt(7)
    caption_style.paragraph_format.space_after = Pt(3)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_border(title._p)
    title.add_run("星河订单服务概要设计待补全文档")
    add_para(document, "本文档记录订单服务的范围、接口、数据和运行要求，供研发、测试及运维人员实施与验收。关键参数在业务事实确认后定稿。")

    add_table(document, ["文档属性", "内容"], [
        ["文档版本", "V1.0"],
        ["负责人", ""],
        ["评审人", ""],
        ["实施日期", ""],
    ], [1.65, 5.25], title="表 1 文档基本信息")

    add_para(document, "系统目标与范围", style="Heading 1")
    add_para(document, "订单服务负责订单创建、订单查询和订单取消。支付结算由独立支付平台处理，现有边界不得因补全操作而改变。")
    add_para(document, "建设目标：待补充")
    add_para(document, "不包含范围：待补充")
    add_styled_placeholder(document, "计划上线窗口：")

    add_para(document, "逻辑架构与部署", style="Heading 1")
    add_para(document, "订单服务通过 HTTP API 接受请求，调用库存适配器完成预留，并将订单状态持久化。API 网关负责统一鉴权。")
    add_para(document, "部署拓扑：待补充")
    add_para(document, "故障隔离策略：待补充")
    add_table(document, ["组件", "职责", "部署位置", "负责人"], [
        ["API 网关", "统一鉴权与路由", "", ""],
        ["订单服务", "", "", "李工"],
        ["库存适配器", "", "", ""],
        ["订单数据库", "订单持久化", "", "数据平台"],
    ], [1.25, 2.2, 2.0, 1.45], title="表 2 核心组件清单")

    add_para(document, "接口设计", style="Heading 1")
    add_para(document, "接口采用 JSON 请求和响应。已确认的接口用途必须保留；新增接口只能来自业务事实清单。")
    add_table(document, ["接口", "方法", "用途", "超时毫秒"], [
        ["/orders", "POST", "", ""],
        ["/orders/{id}", "GET", "查询订单详情", ""],
        ["/orders/{id}/cancel", "POST", "", ""],
    ], [2.15, 0.75, 2.7, 1.3], title="表 3 订单接口清单")
    add_para(document, "请求幂等策略：待补充")

    add_para(document, "数据模型与状态", style="Heading 1")
    add_para(document, "订单主表以 orderId 作为业务标识。字段含义和状态流转以最终审核结果为准。")
    add_table(document, ["字段", "类型", "必填", "说明"], [
        ["orderId", "string", "是", ""],
        ["customerId", "string", "是", ""],
        ["amount", "decimal", "是", ""],
        ["currency", "string", "是", ""],
        ["createdAt", "datetime", "是", "订单保存时间"],
        ["status", "string", "是", ""],
    ], [1.45, 1.15, 0.75, 3.55], title="表 4 订单主表字段")
    add_para(document, "订单状态流转：待补充")

    add_para(document, "异常与风险", style="Heading 1")
    add_para(document, "本节记录实施前必须评估的运行风险。风险项名称属于既有内容，只补全影响与应对措施。")
    risk = add_table(document, ["风险项", "影响", "应对措施"], [
        ["库存超卖", "", ""],
        ["重试产生重复订单", "", ""],
        ["数据库不可用", "", ""],
    ], [1.7, 2.2, 3.0], title="表 5 风险登记")
    # A merged title row makes this table intentionally ragged for the source parser.
    first = risk.add_row()
    risk._tbl.remove(first._tr)
    risk._tbl.insert(0, first._tr)
    merged = first.cells[0].merge(first.cells[2])
    set_cell_text(merged, "核心风险登记", header=True, center=True)
    set_cell_shading(merged, "DDEAF4")
    add_styled_placeholder(document, "评审结论：")

    add_para(document, "非功能需求与验收", style="Heading 1")
    add_para(document, "验收以可复核的测试记录为准。未确认的目标值不得由模型自行推测。")
    add_table(document, ["指标", "目标值", "验收方式"], [
        ["创建接口 P95", "", ""],
        ["服务可用性", "", ""],
        ["重复请求幂等", "", ""],
        ["恢复点目标 RPO", "", ""],
    ], [1.7, 1.65, 3.55], title="表 6 非功能验收指标")

    add_para(document, "上线与回滚", style="Heading 1")
    add_para(document, "上线前须完成接口联调、数据校验与回滚演练，并由项目负责人确认记录。")
    add_para(document, "发布步骤：待补充")
    add_para(document, "回滚触发条件：待补充")
    add_para(document, "回滚操作：待补充")
    add_para(document, "实施日期：待补充")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path,
                        default=Path("target/completion-complex-sample/complex-source.docx"))
    build(parser.parse_args().output)
